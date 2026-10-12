import io
from datetime import date

from django.db.models import Sum, Count, F, Q
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.utils import timezone

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from apps.pasajes.models import Pasaje, PaquetePasaje
from apps.flota.models import Viaje
from apps.caja.models import SesionCaja, Transaccion
from rest_framework.exceptions import ValidationError
from apps.nucleo.permisos import sucursal_permitida

try:
    from apps.encomiendas.models import Encomienda

    ENCOMIENDAS_ENABLED = True
except ImportError:
    ENCOMIENDAS_ENABLED = False


# ── Helpers ───────────────────────────────────────────────────────────────────


def _rango(request):
    hoy = str(timezone.localdate())  # fecha local (Lima), no la UTC
    desde = request.query_params.get("desde") or hoy
    hasta = request.query_params.get("hasta") or hoy
    try:
        f_desde = date.fromisoformat(desde)
        f_hasta = date.fromisoformat(hasta)
    except ValueError:
        raise ValidationError(
            {"error": "Las fechas deben tener formato AAAA-MM-DD (ej. 2026-10-31)."}
        )
    if f_desde > f_hasta:
        raise ValidationError(
            {"error": 'La fecha "desde" no puede ser posterior a "hasta".'}
        )
    return desde, hasta


def _alcance_sucursal(request):
    """
    Sucursal a la que se limita el reporte.
    Cajero: siempre la suya. Administrador/supervisor: la del parámetro ?sucursal= (o todas).
    """
    forzada = sucursal_permitida(request.user)
    if forzada is not None:
        return forzada
    return request.query_params.get("sucursal") or None


def _kw(sid, campo):
    """Filtro opcional: {campo: sid} si hay sucursal, {} si no."""
    return {campo: sid} if sid else {}


def _float(val):
    return float(val or 0)


# ══════════════════════════════════════════════════════════════════════════════
# DASHBOARD
# ══════════════════════════════════════════════════════════════════════════════


class DashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        desde, hasta = _rango(request)
        sid = _alcance_sucursal(request)

        # ── Querysets base ────────────────────────────────────
        pasajes_qs = Pasaje.objects.filter(
            creado_en__date__gte=desde,
            creado_en__date__lte=hasta,
            **_kw(sid, "sucursal_id"),
        ).exclude(estado=Pasaje.Estado.ANULADO)

        viajes_qs = Viaje.objects.filter(
            fecha_salida__date__gte=desde,
            fecha_salida__date__lte=hasta,
        )

        transacciones_qs = Transaccion.objects.filter(
            creado_en__date__gte=desde,
            creado_en__date__lte=hasta,
            **_kw(sid, "sesion__sucursal_id"),
        )

        # ── Paquetes agrupados por día ────────────────────────
        paquetes_dia = (
            PaquetePasaje.objects.filter(
                pasaje__in=pasajes_qs,
                activo=True,
            )
            .annotate(fecha=TruncDate("pasaje__creado_en"))
            .values("fecha")
            .annotate(ingresos_paquetes=Sum("precio_adicional"))
        )
        paquetes_por_dia = {
            str(item["fecha"]): _float(item["ingresos_paquetes"])
            for item in paquetes_dia
        }

        # ── Ventas diarias ────────────────────────────────────
        pasajes_dia = (
            pasajes_qs.annotate(fecha=TruncDate("creado_en"))
            .values("fecha")
            .annotate(pasajes=Count("id"), ingresos_pasajes=Sum("precio"))
        )

        dias = {}
        for p in pasajes_dia:
            f = str(p["fecha"])
            dias[f] = {
                "fecha": f,
                "pasajes": p["pasajes"],
                "ingresos_pasajes": _float(p["ingresos_pasajes"]),
                "ingresos_paquetes": paquetes_por_dia.get(f, 0.0),
                "encomiendas": 0,
                "ingresos_encomiendas": 0.0,
            }

        if ENCOMIENDAS_ENABLED:
            encomiendas_qs = Encomienda.objects.filter(
                activo=True,
                creado_en__date__gte=desde,
                creado_en__date__lte=hasta,
                **_kw(sid, "sucursal_origen_id"),
            ).exclude(
                estado="devuelta"
            )  # devueltas: ya se reembolsaron en caja
            enc_dia = (
                encomiendas_qs.annotate(fecha=TruncDate("creado_en"))
                .values("fecha")
                .annotate(encomiendas=Count("id"), ingresos_encomiendas=Sum("precio"))
            )
            for e in enc_dia:
                f = str(e["fecha"])
                dias.setdefault(
                    f,
                    {
                        "fecha": f,
                        "pasajes": 0,
                        "ingresos_pasajes": 0.0,
                        "ingresos_paquetes": paquetes_por_dia.get(f, 0.0),
                        "encomiendas": 0,
                        "ingresos_encomiendas": 0.0,
                    },
                )
                dias[f]["encomiendas"] = e["encomiendas"]
                dias[f]["ingresos_encomiendas"] = _float(e["ingresos_encomiendas"])
            enc_ingresos = _float(encomiendas_qs.aggregate(t=Sum("precio"))["t"])
            enc_stats = {
                "total": encomiendas_qs.count(),
                # registrada + en_transito = pendientes de entrega
                "en_transito": encomiendas_qs.filter(
                    estado__in=["registrada", "en_transito"]
                ).count(),
                "entregado": encomiendas_qs.filter(estado="entregada").count(),
                "ingresos": enc_ingresos,
            }
        else:
            enc_ingresos = 0.0
            enc_stats = {"total": 0, "en_transito": 0, "entregado": 0, "ingresos": 0.0}

        ventas_diarias = sorted(dias.values(), key=lambda x: x["fecha"])

        # ── Totales pasajes + paquetes ────────────────────────
        total_ingresos_pasajes = _float(pasajes_qs.aggregate(t=Sum("precio"))["t"])

        total_paquetes_ingresos = _float(
            PaquetePasaje.objects.filter(
                pasaje__in=pasajes_qs,
                activo=True,
            ).aggregate(
                t=Sum("precio_adicional")
            )["t"]
        )

        # ── Métodos de pago (pasaje + paquetes + encomiendas por método) ────
        paquetes_por_metodo = {}
        for item in (
            PaquetePasaje.objects.filter(pasaje__in=pasajes_qs, activo=True)
            .values(metodo_pago=F("pasaje__metodo_pago"))
            .annotate(total_paquetes=Sum("precio_adicional"))
        ):
            paquetes_por_metodo[item["metodo_pago"]] = _float(item["total_paquetes"])

        # Encomiendas agrupadas por método de pago
        enc_por_metodo = {}
        if ENCOMIENDAS_ENABLED:
            for item in encomiendas_qs.values("metodo_pago").annotate(
                total_enc=Sum("precio")
            ):
                enc_por_metodo[item["metodo_pago"]] = _float(item["total_enc"])

        # Unir todos los métodos (pasajes + paquetes + encomiendas)
        todos_metodos = set(
            list(pasajes_qs.values_list("metodo_pago", flat=True).distinct())
            + list(enc_por_metodo.keys())
        )

        total_base_mp = (
            total_ingresos_pasajes + total_paquetes_ingresos + enc_ingresos
        ) or 1
        metodos_pago = []
        for metodo in todos_metodos:
            total_pasaje = _float(
                pasajes_qs.filter(metodo_pago=metodo).aggregate(t=Sum("precio"))["t"]
            )
            total_paquete = paquetes_por_metodo.get(metodo, 0.0)
            total_enc_mp = enc_por_metodo.get(metodo, 0.0)
            total_mp = total_pasaje + total_paquete + total_enc_mp
            if total_mp > 0:
                metodos_pago.append(
                    {
                        "metodo": metodo,
                        "total": total_mp,
                        "cantidad": pasajes_qs.filter(metodo_pago=metodo).count(),
                        "porcentaje": round(total_mp / total_base_mp * 100, 1),
                    }
                )

        # ── Rutas top ─────────────────────────────────────────
        rutas_top = []
        for item in (
            pasajes_qs.values(
                origen=F("viaje__ruta__origen__nombre"),
                destino=F("viaje__ruta__destino__nombre"),
            )
            .annotate(total=Count("id"))
            .order_by("-total")[:5]
        ):
            rutas_top.append(
                {
                    "ruta": f"{item['origen']} → {item['destino']}",
                    "total": item["total"],
                }
            )

        # ── Viajes por estado ─────────────────────────────────
        viajes_estado = {}
        for item in viajes_qs.values("estado").annotate(total=Count("id")):
            viajes_estado[item["estado"]] = item["total"]

        # ── Caja ──────────────────────────────────────────────
        # FIX: una sola asignación limpia con _float(), restando anulaciones
        ingresos_caja_bruto = _float(
            transacciones_qs.filter(
                tipo__in=["venta_pasaje", "venta_encomienda"]
            ).aggregate(t=Sum("monto"))["t"]
        )
        anulaciones_caja = _float(
            transacciones_qs.filter(tipo="anulacion").aggregate(t=Sum("monto"))["t"]
        )
        ingresos_caja = ingresos_caja_bruto - anulaciones_caja

        egresos_caja = _float(
            transacciones_qs.filter(tipo="gasto").aggregate(t=Sum("monto"))["t"]
        )
        total_sesiones = SesionCaja.objects.filter(
            fecha_apertura__date__gte=desde,
            fecha_apertura__date__lte=hasta,
            **_kw(sid, "sucursal_id"),
        ).count()

        # ── Total ventas = pasajes + paquetes + encomiendas ───
        total_ventas = total_ingresos_pasajes + total_paquetes_ingresos + enc_ingresos

        return Response(
            {
                "total_ventas": total_ventas,
                "total_pasajes": pasajes_qs.count(),
                "total_viajes": viajes_qs.count(),
                "ventas_diarias": ventas_diarias,
                "metodos_pago": metodos_pago,
                "rutas_top": rutas_top,
                "viajes_estado": viajes_estado,
                "encomiendas": enc_stats,
                "caja": {
                    "total_ingresos": ingresos_caja,  # neto: ventas - anulaciones
                    "total_egresos": egresos_caja,
                    "total_sesiones": total_sesiones,
                },
            }
        )


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE VENTAS
# ══════════════════════════════════════════════════════════════════════════════


class ReporteVentasView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        desde, hasta = _rango(request)
        sucursal_id = _alcance_sucursal(request)

        filtros_pasaje = Q(creado_en__date__gte=desde, creado_en__date__lte=hasta)
        filtros_encomienda = Q(creado_en__date__gte=desde, creado_en__date__lte=hasta)

        if sucursal_id:
            filtros_pasaje &= Q(sucursal_id=sucursal_id)
            filtros_encomienda &= Q(sucursal_origen_id=sucursal_id)

        pasajes_qs_activos = Pasaje.objects.filter(
            filtros_pasaje,
            estado__in=[Pasaje.Estado.EMITIDO, Pasaje.Estado.EMBARCADO],
        )

        pasajes = pasajes_qs_activos.aggregate(
            total=Count("id"), ingresos=Sum("precio")
        )

        paquetes_ingresos = _float(
            PaquetePasaje.objects.filter(
                pasaje__in=pasajes_qs_activos,
                activo=True,
            ).aggregate(t=Sum("precio_adicional"))["t"]
        )

        anulaciones = Pasaje.objects.filter(
            filtros_pasaje,
            estado=Pasaje.Estado.ANULADO,
        ).aggregate(total=Count("id"), monto=Sum("precio"))

        enc_data = {"total": 0, "ingresos": 0.0}
        if ENCOMIENDAS_ENABLED:
            enc = (
                Encomienda.objects.filter(filtros_encomienda, activo=True)
                .exclude(estado="devuelta")
                .aggregate(total=Count("id"), ingresos=Sum("precio"))
            )
            enc_data = {"total": enc["total"] or 0, "ingresos": _float(enc["ingresos"])}

        ingresos_pasajes = _float(pasajes["ingresos"])

        return Response(
            {
                "pasajes": {
                    "total": pasajes["total"] or 0,
                    "ingresos": ingresos_pasajes,
                },
                "paquetes": {"ingresos": paquetes_ingresos},
                "encomiendas": enc_data,
                "anulaciones": {
                    "total": anulaciones["total"] or 0,
                    "monto": _float(anulaciones["monto"]),
                },
                "total_ingresos": ingresos_pasajes
                + paquetes_ingresos
                + enc_data["ingresos"],
            }
        )


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE CAJA
# ══════════════════════════════════════════════════════════════════════════════


class ReporteCajaView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        desde, hasta = _rango(request)
        sucursal_id = _alcance_sucursal(request)

        filtros = Q(
            fecha_apertura__date__gte=desde,
            fecha_apertura__date__lte=hasta,
        )
        if sucursal_id:
            filtros &= Q(sucursal_id=sucursal_id)

        sesiones = SesionCaja.objects.filter(filtros).select_related(
            "cajero", "sucursal"
        )

        resumen = []
        for s in sesiones:
            resumen.append(
                {
                    "cajero": getattr(s.cajero, "nombre_completo", str(s.cajero)),
                    "sucursal": s.sucursal.nombre if s.sucursal else "—",
                    "monto_inicial": _float(s.monto_inicial),
                    "total_ingresos": _float(s.total_ingresos),
                    "total_egresos": _float(s.total_egresos),
                    "saldo_esperado": _float(s.saldo_esperado),
                    "monto_final": (
                        _float(s.monto_final) if s.monto_final is not None else None
                    ),
                    "diferencia": (
                        round(_float(s.monto_final) - _float(s.saldo_esperado), 2)
                        if s.monto_final is not None
                        else None
                    ),
                    "estado": s.estado,
                    "fecha_apertura": s.fecha_apertura,
                    "fecha_cierre": s.fecha_cierre,
                }
            )

        return Response({"desde": desde, "hasta": hasta, "sesiones": resumen})


# ══════════════════════════════════════════════════════════════════════════════
# EXPORTAR EXCEL
# 3 hojas: Resumen Total | Pasajes y Bultos | Encomiendas
# Reglas de negocio:
#   - Pasajes ANULADOS excluidos → sus paquetes activo=False también excluidos
#   - Encomiendas DEVUELTAS excluidas de ingresos
# ══════════════════════════════════════════════════════════════════════════════


class ExportarExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        desde, hasta = _rango(request)
        sid = _alcance_sucursal(request)

        # ── Querysets base con reglas de negocio ──────────────
        # Pasajes activos (no anulados) → paquetes con activo=True automático
        pasajes_qs = (
            Pasaje.objects.filter(
                creado_en__date__gte=desde,
                creado_en__date__lte=hasta,
                **_kw(sid, "sucursal_id"),
            )
            .exclude(estado=Pasaje.Estado.ANULADO)
            .select_related(
                "pasajero", "viaje__ruta__origen", "viaje__ruta__destino", "vendido_por"
            )
            .prefetch_related("paquetes")
            .order_by("creado_en")
        )
        pasajes_list = list(pasajes_qs)

        # Encomiendas activas (excluir devueltas de ingresos)
        enc_qs = None
        enc_list = []
        if ENCOMIENDAS_ENABLED:
            enc_qs = (
                Encomienda.objects.filter(
                    activo=True,
                    creado_en__date__gte=desde,
                    creado_en__date__lte=hasta,
                    **_kw(sid, "sucursal_origen_id"),
                )
                .exclude(estado="devuelta")  # devueltas no cuentan como ingreso
                .select_related("registrado_por", "sucursal_origen", "viaje")
                .order_by("creado_en")
            )
            enc_list = list(enc_qs)

        # ── Totales globales ──────────────────────────────────
        total_pasajes_monto = sum(float(p.precio or 0) for p in pasajes_list)
        total_paquetes_monto = sum(float(p.total_paquetes or 0) for p in pasajes_list)
        total_enc_monto = sum(float(e.precio or 0) for e in enc_list)
        gran_total = total_pasajes_monto + total_paquetes_monto + total_enc_monto

        # ── Estilos comunes ───────────────────────────────────
        wb = Workbook()

        AZUL = "FF1E3A5F"
        NARANJA = "FFEA580C"
        NARANJA_CLARO = "FFFDE8D8"
        VERDE_CLARO = "FFE6F4EA"
        AZUL_CLARO = "FFEFF6FF"
        BLANCO = "FFFFFFFF"
        GRIS = "FFF3F4F6"
        GRIS_MED = "FFE5E7EB"

        h_font = Font(name="Calibri", bold=True, color=BLANCO, size=11)
        title_font = Font(name="Calibri", bold=True, size=14, color=AZUL)
        sub_font = Font(name="Calibri", bold=True, size=10, color=AZUL)
        h_fill = PatternFill("solid", fgColor=AZUL)
        alt_fill = PatternFill("solid", fgColor=GRIS)
        total_fill = PatternFill("solid", fgColor=NARANJA)

        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center")
        right = Alignment(horizontal="right", vertical="center")
        thin = Side(style="thin", color="FFD1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        METODO_LABEL = {
            "cash": "Efectivo",
            "yape": "Yape",
            "plin": "Plin",
            "card": "Tarjeta",
        }
        ESTADO_ENC = {
            "registrada": "Registrada",
            "en_transito": "En Tránsito",
            "entregada": "Entregada",
            "devuelta": "Devuelta",
        }

        def _header(ws, row, col, val):
            c = ws.cell(row=row, column=col, value=val)
            c.font = h_font
            c.fill = h_fill
            c.alignment = center
            c.border = border
            return c

        def _cell(ws, row, col, val, aln=None, bold=False, fill=None, color=None):
            c = ws.cell(row=row, column=col, value=val)
            c.font = Font(
                name="Calibri", size=10, bold=bold, color=color if color else "00000000"
            )
            c.alignment = aln or left
            c.border = border
            if fill:
                c.fill = fill
            return c

        def _titulo(ws, titulo, cols):
            ws.merge_cells(f"A1:{chr(64+cols)}1")
            t = ws["A1"]
            t.value = titulo
            t.font = title_font
            t.alignment = center
            t.fill = PatternFill("solid", fgColor=NARANJA_CLARO)
            ws.row_dimensions[1].height = 28

        def _fila_total(ws, row, cols_vals, total_cols):
            """cols_vals: [(col, val), ...]. total_cols: número de columnas a fusionar al inicio"""
            ws.merge_cells(f"A{row}:{chr(64+total_cols)}{row}")
            for col, val in cols_vals:
                c = ws.cell(row=row, column=col, value=val)
                c.font = Font(name="Calibri", bold=True, size=11, color=BLANCO)
                c.fill = total_fill
                c.alignment = right if col > total_cols else center
                c.border = border

        # ══════════════════════════════════════════════════════
        # HOJA 1: RESUMEN TOTAL
        # ══════════════════════════════════════════════════════
        ws1 = wb.active
        ws1.title = "Resumen Total"
        ws1.sheet_view.showGridLines = False

        _titulo(ws1, f"RESUMEN TOTAL — {desde} al {hasta}", 4)

        # ── Bloque KPIs ──────────────────────────────────────
        ws1.row_dimensions[3].height = 20
        ws1["A3"] = "RESUMEN FINANCIERO"
        ws1["A3"].font = sub_font

        kpi_headers = ["Concepto", "Cantidad", "Ingresos (S/.)", "Observación"]
        ws1.row_dimensions[4].height = 22
        for col, h in enumerate(kpi_headers, 1):
            _header(ws1, 4, col, h)

        kpi_rows = [
            (
                "Pasajes vendidos",
                len(pasajes_list),
                total_pasajes_monto,
                "Sin anulados",
            ),
            (
                "Bultos / Paquetes",
                sum(p.paquetes.filter(activo=True).count() for p in pasajes_list),
                total_paquetes_monto,
                "Activos, de pasajes vigentes",
            ),
            ("Encomiendas", len(enc_list), total_enc_monto, "Sin devueltas"),
            ("GRAN TOTAL", "", gran_total, ""),
        ]
        for i, (concepto, cantidad, monto, obs) in enumerate(kpi_rows, 5):
            is_total = concepto == "GRAN TOTAL"
            fill_row = (
                PatternFill("solid", fgColor=AZUL_CLARO) if not is_total else total_fill
            )
            color = BLANCO if is_total else "00000000"
            bold = is_total
            _cell(ws1, i, 1, concepto, bold=bold, fill=fill_row, color=color)
            _cell(ws1, i, 2, cantidad, right, bold=bold, fill=fill_row, color=color)
            _cell(
                ws1,
                i,
                3,
                round(monto, 2) if monto != "" else "",
                right,
                bold=bold,
                fill=fill_row,
                color=color,
            )
            _cell(ws1, i, 4, obs, fill=fill_row, color=color)
            ws1.row_dimensions[i].height = 18

        # ── Bloque métodos de pago ────────────────────────────
        row_mp = len(kpi_rows) + 6
        ws1.row_dimensions[row_mp].height = 20
        ws1[f"A{row_mp}"] = "MÉTODOS DE PAGO"
        ws1[f"A{row_mp}"].font = sub_font

        mp_headers = [
            "Método",
            "Pasajes (S/.)",
            "Paquetes (S/.)",
            "Encomiendas (S/.)",
            "Total (S/.)",
        ]
        row_mp += 1
        ws1.row_dimensions[row_mp].height = 22
        for col, h in enumerate(mp_headers, 1):
            _header(ws1, row_mp, col, h)

        # Calcular por método
        metodos = {}
        for p in pasajes_list:
            m = p.metodo_pago
            if m not in metodos:
                metodos[m] = {"pasajes": 0.0, "paquetes": 0.0, "encomiendas": 0.0}
            metodos[m]["pasajes"] += float(p.precio or 0)
            metodos[m]["paquetes"] += float(p.total_paquetes or 0)
        for e in enc_list:
            m = e.metodo_pago
            if m not in metodos:
                metodos[m] = {"pasajes": 0.0, "paquetes": 0.0, "encomiendas": 0.0}
            metodos[m]["encomiendas"] += float(e.precio or 0)

        tot_mp = {"pasajes": 0.0, "paquetes": 0.0, "encomiendas": 0.0}
        for i, (metodo, vals) in enumerate(sorted(metodos.items()), 1):
            r = row_mp + i
            fill = alt_fill if i % 2 == 0 else None
            total_m = vals["pasajes"] + vals["paquetes"] + vals["encomiendas"]
            _cell(ws1, r, 1, METODO_LABEL.get(metodo, metodo), fill=fill)
            _cell(ws1, r, 2, round(vals["pasajes"], 2), right, fill=fill)
            _cell(ws1, r, 3, round(vals["paquetes"], 2), right, fill=fill)
            _cell(ws1, r, 4, round(vals["encomiendas"], 2), right, fill=fill)
            _cell(ws1, r, 5, round(total_m, 2), right, True, fill=fill)
            ws1.row_dimensions[r].height = 18
            tot_mp["pasajes"] += vals["pasajes"]
            tot_mp["paquetes"] += vals["paquetes"]
            tot_mp["encomiendas"] += vals["encomiendas"]

        tr_mp = row_mp + len(metodos) + 1
        _fila_total(
            ws1,
            tr_mp,
            [
                (1, "TOTAL"),
                (2, round(tot_mp["pasajes"], 2)),
                (3, round(tot_mp["paquetes"], 2)),
                (4, round(tot_mp["encomiendas"], 2)),
                (5, round(gran_total, 2)),
            ],
            1,
        )

        for col, w in zip("ABCDE", [28, 16, 16, 18, 14]):
            ws1.column_dimensions[col].width = w

        # ══════════════════════════════════════════════════════
        # HOJA 2: PASAJES Y BULTOS
        # ══════════════════════════════════════════════════════
        ws2 = wb.create_sheet("Pasajes y Bultos")
        ws2.sheet_view.showGridLines = False

        _titulo(ws2, f"PASAJES Y BULTOS — {desde} al {hasta}", 11)

        headers2 = [
            "#",
            "Código",
            "Fecha",
            "Pasajero",
            "DNI",
            "Ruta",
            "Asiento",
            "Método",
            "Pasaje (S/.)",
            "Bultos (S/.)",
            "Total (S/.)",
        ]
        ws2.row_dimensions[2].height = 22
        for col, h in enumerate(headers2, 1):
            _header(ws2, 2, col, h)

        total_p = 0.0
        total_q = 0.0
        for i, p in enumerate(pasajes_list, 1):
            row = i + 2
            ruta = (
                f"{p.viaje.ruta.origen.nombre} → {p.viaje.ruta.destino.nombre}"
                if p.viaje and p.viaje.ruta
                else "—"
            )
            fill = alt_fill if i % 2 == 0 else None
            precio = float(p.precio or 0)
            bultos = float(p.total_paquetes or 0)
            _cell(ws2, row, 1, i, center, fill=fill)
            _cell(ws2, row, 2, str(p.codigo or "")[:8].upper(), center, fill=fill)
            _cell(
                ws2, row, 3, p.creado_en.strftime("%d/%m/%Y %H:%M"), center, fill=fill
            )
            _cell(
                ws2,
                row,
                4,
                f"{p.pasajero.nombres} {p.pasajero.apellidos}" if p.pasajero else "—",
                fill=fill,
            )
            _cell(ws2, row, 5, p.pasajero.dni if p.pasajero else "—", center, fill=fill)
            _cell(ws2, row, 6, ruta, fill=fill)
            _cell(ws2, row, 7, p.asiento, center, fill=fill)
            _cell(
                ws2,
                row,
                8,
                METODO_LABEL.get(p.metodo_pago, p.metodo_pago),
                center,
                fill=fill,
            )
            _cell(ws2, row, 9, precio, right, fill=fill)
            _cell(ws2, row, 10, bultos, right, fill=fill)
            _cell(ws2, row, 11, round(precio + bultos, 2), right, True, fill=fill)
            ws2.row_dimensions[row].height = 18
            total_p += precio
            total_q += bultos

        _fila_total(
            ws2,
            len(pasajes_list) + 3,
            [
                (1, "TOTAL"),
                (9, round(total_p, 2)),
                (10, round(total_q, 2)),
                (11, round(total_p + total_q, 2)),
            ],
            8,
        )

        for col, w in zip("ABCDEFGHIJK", [5, 12, 18, 28, 12, 34, 9, 12, 13, 13, 13]):
            ws2.column_dimensions[col].width = w

        # ══════════════════════════════════════════════════════
        # HOJA 3: ENCOMIENDAS
        # ══════════════════════════════════════════════════════
        ws3 = wb.create_sheet("Encomiendas")
        ws3.sheet_view.showGridLines = False

        _titulo(ws3, f"ENCOMIENDAS — {desde} al {hasta}", 10)

        headers3 = [
            "#",
            "Código",
            "Fecha",
            "Remitente",
            "Destinatario",
            "Descripción",
            "Peso (kg)",
            "Método",
            "Precio (S/.)",
            "Estado",
        ]
        ws3.row_dimensions[2].height = 22
        for col, h in enumerate(headers3, 1):
            _header(ws3, 2, col, h)

        ESTADO_COLOR = {
            "registrada": "FF6B7280",  # gris
            "en_transito": "FF2563EB",  # azul
            "entregada": "FF16A34A",  # verde
            "devuelta": "FFDC2626",  # rojo
        }

        total_enc = 0.0
        for i, e in enumerate(enc_list, 1):
            row = i + 2
            fill = alt_fill if i % 2 == 0 else None
            precio = float(e.precio or 0)
            estado_val = e.estado
            _cell(ws3, row, 1, i, center, fill=fill)
            _cell(ws3, row, 2, str(e.codigo)[:8].upper(), center, fill=fill)
            _cell(
                ws3, row, 3, e.creado_en.strftime("%d/%m/%Y %H:%M"), center, fill=fill
            )
            _cell(ws3, row, 4, e.remitente_nombre, fill=fill)
            _cell(ws3, row, 5, e.destinatario_nombre, fill=fill)
            _cell(
                ws3,
                row,
                6,
                e.descripcion[:40] + ("…" if len(e.descripcion) > 40 else ""),
                fill=fill,
            )
            _cell(
                ws3, row, 7, float(e.peso_kg) if e.peso_kg else "—", center, fill=fill
            )
            _cell(
                ws3,
                row,
                8,
                METODO_LABEL.get(e.metodo_pago, e.metodo_pago),
                center,
                fill=fill,
            )
            _cell(ws3, row, 9, precio, right, True, fill=fill)
            # Estado con color
            c_estado = ws3.cell(
                row=row, column=10, value=ESTADO_ENC.get(estado_val, estado_val)
            )
            c_estado.font = Font(
                name="Calibri",
                size=10,
                bold=True,
                color=ESTADO_COLOR.get(estado_val, "00000000"),
            )
            c_estado.alignment = center
            c_estado.border = border
            if fill:
                c_estado.fill = fill
            ws3.row_dimensions[row].height = 18
            total_enc += precio

        _fila_total(
            ws3,
            len(enc_list) + 3,
            [
                (1, "TOTAL"),
                (9, round(total_enc, 2)),
            ],
            8,
        )

        for col, w in zip("ABCDEFGHIJ", [5, 12, 18, 28, 28, 30, 10, 12, 14, 14]):
            ws3.column_dimensions[col].width = w

        # ── Guardar ───────────────────────────────────────────
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"reporte_{desde}_{hasta}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


# ══════════════════════════════════════════════════════════════════════════════
# EXPORTAR PDF — A4 vertical, márgenes 1.5cm laterales
# Mismas reglas: pasajes no anulados, encomiendas no devueltas
# ══════════════════════════════════════════════════════════════════════════════


class ExportarPDFView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        desde, hasta = _rango(request)
        sid = _alcance_sucursal(request)

        # ── Querysets con reglas de negocio ───────────────────
        pasajes_qs = (
            Pasaje.objects.filter(
                creado_en__date__gte=desde,
                creado_en__date__lte=hasta,
                **_kw(sid, "sucursal_id"),
            )
            .exclude(estado=Pasaje.Estado.ANULADO)
            .select_related("pasajero", "viaje__ruta__origen", "viaje__ruta__destino")
            .prefetch_related("paquetes")
            .order_by("creado_en")
        )
        pasajes_list = list(pasajes_qs)

        enc_list = []
        if ENCOMIENDAS_ENABLED:
            enc_list = list(
                Encomienda.objects.filter(
                    activo=True,
                    creado_en__date__gte=desde,
                    creado_en__date__lte=hasta,
                    **_kw(sid, "sucursal_origen_id"),
                )
                .exclude(estado="devuelta")
                .order_by("creado_en")
            )

        # ── Totales ───────────────────────────────────────────
        total_pasajes_monto = sum(float(p.precio or 0) for p in pasajes_list)
        total_paquetes_monto = sum(float(p.total_paquetes or 0) for p in pasajes_list)
        total_enc_monto = sum(float(e.precio or 0) for e in enc_list)
        gran_total = total_pasajes_monto + total_paquetes_monto + total_enc_monto

        # ── Documento PDF A4 vertical ─────────────────────────
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=1.5 * cm,
            leftMargin=1.5 * cm,
            topMargin=2 * cm,
            bottomMargin=1.5 * cm,
        )

        # Ancho útil = 21cm - 3cm márgenes = 18cm
        PAGE_W = 18 * cm

        styles = getSampleStyleSheet()
        C_AZUL = colors.HexColor("#1E3A5F")
        C_NARANJA = colors.HexColor("#EA580C")
        C_GRIS = colors.HexColor("#F3F4F6")
        C_BLANCO = colors.white
        C_VERDE = colors.HexColor("#16A34A")
        C_ROJO = colors.HexColor("#DC2626")

        titulo_style = ParagraphStyle(
            "titulo",
            parent=styles["Title"],
            fontSize=15,
            textColor=C_AZUL,
            spaceAfter=2,
        )
        sub_style = ParagraphStyle(
            "sub",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#6B7280"),
            spaceAfter=10,
        )
        seccion_style = ParagraphStyle(
            "seccion",
            parent=styles["Normal"],
            fontSize=11,
            textColor=C_AZUL,
            fontName="Helvetica-Bold",
            spaceAfter=6,
            spaceBefore=12,
        )

        METODO_LABEL = {
            "cash": "Efectivo",
            "yape": "Yape",
            "plin": "Plin",
            "card": "Tarjeta",
        }

        def _ts_header():
            return [
                ("BACKGROUND", (0, 0), (-1, 0), C_AZUL),
                ("TEXTCOLOR", (0, 0), (-1, 0), C_BLANCO),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -2), 7.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [C_BLANCO, C_GRIS]),
                ("BACKGROUND", (0, -1), (-1, -1), C_NARANJA),
                ("TEXTCOLOR", (0, -1), (-1, -1), C_BLANCO),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, -1), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D1D5DB")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]

        elements = []

        # ── Encabezado ────────────────────────────────────────
        elements.append(Paragraph("WJL Transportes — Reporte General", titulo_style))
        elements.append(
            Paragraph(
                f'Período: {desde} al {hasta}  |  Generado: {timezone.now().strftime("%d/%m/%Y %H:%M")}',
                sub_style,
            )
        )

        # ── KPIs resumen ──────────────────────────────────────
        kpi_data = [
            ["Pasajes", "Bultos/Paquetes", "Encomiendas", "Gran Total"],
            [
                f"{len(pasajes_list)} — S/. {total_pasajes_monto:,.2f}",
                f"S/. {total_paquetes_monto:,.2f}",
                f"{len(enc_list)} — S/. {total_enc_monto:,.2f}",
                f"S/. {gran_total:,.2f}",
            ],
        ]
        kpi_w = PAGE_W / 4
        kpi_t = Table(kpi_data, colWidths=[kpi_w] * 4)
        kpi_t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), C_AZUL),
                    ("TEXTCOLOR", (0, 0), (-1, 0), C_BLANCO),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EFF6FF")),
                    ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 1), (-1, 1), 9),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    # Gran total destacado
                    ("BACKGROUND", (3, 1), (3, 1), colors.HexColor("#FFF7ED")),
                    ("TEXTCOLOR", (3, 1), (3, 1), C_NARANJA),
                ]
            )
        )
        elements.append(kpi_t)
        elements.append(Spacer(1, 0.4 * cm))

        # ── Sección Pasajes ───────────────────────────────────
        elements.append(Paragraph("Pasajes vendidos", seccion_style))

        # Columnas A4 vertical: #, Código, Fecha, Pasajero, Ruta, Asiento, Método, Pasaje, Bultos, Total
        col_w2 = [
            0.5 * cm,
            1.6 * cm,
            2.6 * cm,
            3.8 * cm,
            3.5 * cm,
            1.2 * cm,
            1.6 * cm,
            1.6 * cm,
            1.6 * cm,
            1.5 * cm,
        ]
        rows_p = [
            [
                "#",
                "Código",
                "Fecha",
                "Pasajero",
                "Ruta",
                "Asiento",
                "Método",
                "Pasaje\n(S/.)",
                "Bultos\n(S/.)",
                "Total\n(S/.)",
            ]
        ]
        for i, p in enumerate(pasajes_list, 1):
            ruta = (
                f"{p.viaje.ruta.origen.nombre}→{p.viaje.ruta.destino.nombre}"
                if p.viaje and p.viaje.ruta
                else "—"
            )
            precio_p = float(p.precio or 0)
            precio_q = float(p.total_paquetes or 0)
            rows_p.append(
                [
                    str(i),
                    str(p.codigo or "")[:8].upper(),
                    p.creado_en.strftime("%d/%m/%Y\n%H:%M"),
                    (
                        f"{p.pasajero.nombres} {p.pasajero.apellidos}"
                        if p.pasajero
                        else "—"
                    ),
                    ruta,
                    str(p.asiento),
                    METODO_LABEL.get(p.metodo_pago, p.metodo_pago or "—"),
                    f"{precio_p:,.2f}",
                    f"{precio_q:,.2f}",
                    f"{precio_p + precio_q:,.2f}",
                ]
            )
        rows_p.append(
            [
                "",
                "",
                "",
                "",
                "",
                "",
                "TOTAL",
                f"S/. {total_pasajes_monto:,.2f}",
                f"S/. {total_paquetes_monto:,.2f}",
                f"S/. {total_pasajes_monto + total_paquetes_monto:,.2f}",
            ]
        )

        ts_p = _ts_header() + [
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("ALIGN", (-3, 0), (-1, -1), "RIGHT"),
            ("ALIGN", (5, 1), (6, -1), "CENTER"),
        ]
        t_p = Table(rows_p, colWidths=col_w2, repeatRows=1)
        t_p.setStyle(TableStyle(ts_p))
        elements.append(t_p)
        elements.append(Spacer(1, 0.4 * cm))

        # ── Sección Encomiendas ───────────────────────────────
        if enc_list:
            elements.append(Paragraph("Encomiendas", seccion_style))

            # Columnas: #, Código, Fecha, Remitente, Destinatario, Descripción, Método, Precio, Estado
            col_w3 = [
                0.5 * cm,
                1.6 * cm,
                2.3 * cm,
                3.2 * cm,
                3.2 * cm,
                3.0 * cm,
                1.5 * cm,
                1.5 * cm,
                1.2 * cm,
            ]
            rows_e = [
                [
                    "#",
                    "Código",
                    "Fecha",
                    "Remitente",
                    "Destinatario",
                    "Descripción",
                    "Método",
                    "Precio\n(S/.)",
                    "Estado",
                ]
            ]
            ESTADO_ENC = {
                "registrada": "Regist.",
                "en_transito": "Tránsito",
                "entregada": "Entregada",
                "devuelta": "Devuelta",
            }
            for i, e in enumerate(enc_list, 1):
                rows_e.append(
                    [
                        str(i),
                        str(e.codigo)[:8].upper(),
                        e.creado_en.strftime("%d/%m/%Y\n%H:%M"),
                        e.remitente_nombre[:22],
                        e.destinatario_nombre[:22],
                        (
                            (e.descripcion[:28] + "…")
                            if len(e.descripcion) > 28
                            else e.descripcion
                        ),
                        METODO_LABEL.get(e.metodo_pago, e.metodo_pago),
                        f"{float(e.precio or 0):,.2f}",
                        ESTADO_ENC.get(e.estado, e.estado),
                    ]
                )
            rows_e.append(
                ["", "", "", "", "", "", "TOTAL", f"S/. {total_enc_monto:,.2f}", ""]
            )

            ts_e = _ts_header() + [
                ("ALIGN", (0, 1), (0, -1), "CENTER"),
                ("ALIGN", (-2, 0), (-2, -1), "RIGHT"),
                ("ALIGN", (-1, 0), (-1, -1), "CENTER"),
                ("ALIGN", (2, 1), (2, -1), "CENTER"),
            ]
            t_e = Table(rows_e, colWidths=col_w3, repeatRows=1)
            t_e.setStyle(TableStyle(ts_e))
            elements.append(t_e)

        # ── Footer ────────────────────────────────────────────
        def footer(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#9CA3AF"))
            canvas.drawString(
                1.5 * cm, 1 * cm, f"WJL Transportes — Reporte {desde} al {hasta}"
            )
            canvas.drawRightString(A4[0] - 1.5 * cm, 1 * cm, f"Página {doc.page}")
            canvas.restoreState()

        doc.build(elements, onFirstPage=footer, onLaterPages=footer)
        buffer.seek(0)

        filename = f"reporte_{desde}_{hasta}.pdf"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
